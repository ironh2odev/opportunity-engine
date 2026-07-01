import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from app.main import app


FIXTURE_DIR = Path(__file__).resolve().parent / "fixtures"


def _read_fixture(name: str) -> str:
    return (FIXTURE_DIR / name).read_text(encoding="utf-8")


class PersonalLeadsApiTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path = os.path.join(self.temp_dir.name, "personal_test.db")
        os.environ["OE_PERSONAL_DB_PATH"] = self.db_path
        self.old_openai_key = os.environ.get("OPENAI_API_KEY")
        if "OPENAI_API_KEY" in os.environ:
            del os.environ["OPENAI_API_KEY"]
        self.client = TestClient(app)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()
        if "OE_PERSONAL_DB_PATH" in os.environ:
            del os.environ["OE_PERSONAL_DB_PATH"]
        if self.old_openai_key is not None:
            os.environ["OPENAI_API_KEY"] = self.old_openai_key
        elif "OPENAI_API_KEY" in os.environ:
            del os.environ["OPENAI_API_KEY"]

    def _create_lead(self) -> str:
        response = self.client.post(
            "/personal/leads",
            json={
                "name": "Avery Pike",
                "role": "Founder",
                "organisation": "Northfield Systems",
                "source": "manual",
                "opportunity_type": "client",
                "relationship_strength": "cold",
                "status": "new",
                "fit_score": 7,
                "priority": "high",
                "problem_observed": "Slow lead response workflow",
                "why_relevant": "Clear fit with workflow automation offer",
                "suggested_angle": "Start with lightweight process audit",
                "notes": "First pass note",
                "tags": ["saas", "revops"],
                "next_action": "prepare connection request",
            },
        )
        self.assertEqual(response.status_code, 200)
        return response.json()["id"]

    def test_create_lead(self) -> None:
        lead_id = self._create_lead()
        get_response = self.client.get(f"/personal/leads/{lead_id}")
        self.assertEqual(get_response.status_code, 200)
        self.assertEqual(get_response.json()["name"], "Avery Pike")

    def test_update_lead(self) -> None:
        lead_id = self._create_lead()
        update_response = self.client.patch(
            f"/personal/leads/{lead_id}",
            json={"status": "reviewed", "priority": "medium", "notes": "Updated note"},
        )
        self.assertEqual(update_response.status_code, 200)
        self.assertEqual(update_response.json()["status"], "reviewed")
        self.assertEqual(update_response.json()["priority"], "medium")

    def test_import_csv(self) -> None:
        csv_text = """name,role,organisation,organisationWebsite,linkedinUrl,email,location,source,opportunityType,relationshipStrength,problemObserved,whyRelevant,suggestedAngle,notes,tags,nextAction,followUpDate
Riley North,Growth Lead,Summit Relay,https://summit.example,https://linkedin.example/riley,riley@example.com,Austin,csv,client,warm,Lead qualification lag,Good fit with advisory ops,Use roadmap-first pitch,Note one,revops|b2b,review website,2026-07-10
"""
        import_response = self.client.post(
            "/personal/leads/import-csv",
            json={"csv_text": csv_text},
        )
        self.assertEqual(import_response.status_code, 200)
        self.assertEqual(len(import_response.json()), 1)

    def test_invalid_csv_rejected(self) -> None:
        csv_text = """name,role,organisation,organisationWebsite,linkedinUrl,email,location,source,opportunityType,relationshipStrength,problemObserved,whyRelevant,suggestedAngle,notes,tags,nextAction,followUpDate
Jordan Vale,Founder,Cobalt Ridge,https://cobalt.example,https://linkedin.example/jordan,jordan@example.com,Remote,csv,unknown_type,warm,Ops pain,Interesting,Angle,Note,t1,prepare outreach,
"""
        import_response = self.client.post(
            "/personal/leads/import-csv",
            json={"csv_text": csv_text},
        )
        self.assertEqual(import_response.status_code, 400)
        self.assertIn("Invalid CSV row", import_response.json()["detail"])

    def test_create_rule_action_from_lead(self) -> None:
        lead_id = self._create_lead()
        action_response = self.client.post(f"/personal/leads/{lead_id}/create-rule-action")
        self.assertEqual(action_response.status_code, 200)
        action = action_response.json()
        self.assertEqual(action["status"], "suggested")
        self.assertIn("approval_required", action)
        self.assertEqual(action["source_lead_id"], lead_id)

    def test_linked_actions_fetch_for_lead(self) -> None:
        lead_id = self._create_lead()
        self.client.post(f"/personal/leads/{lead_id}/create-rule-action")

        list_response = self.client.get(f"/personal/leads/{lead_id}/rule-actions")
        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(len(list_response.json()), 1)
        self.assertEqual(list_response.json()[0]["source_lead_id"], lead_id)

    def test_linked_actions_survive_reinitialization(self) -> None:
        lead_id = self._create_lead()
        self.client.post(f"/personal/leads/{lead_id}/create-rule-action")

        new_client = TestClient(app)
        list_response = new_client.get(f"/personal/leads/{lead_id}/rule-actions")
        self.assertEqual(list_response.status_code, 200)
        self.assertGreaterEqual(len(list_response.json()), 1)

    def test_invalid_lead_id_returns_clear_error(self) -> None:
        create_response = self.client.post("/personal/leads/lead_unknown/create-rule-action")
        self.assertEqual(create_response.status_code, 404)
        self.assertIn("Lead not found", create_response.json()["detail"])

        list_response = self.client.get("/personal/leads/lead_unknown/rule-actions")
        self.assertEqual(list_response.status_code, 404)
        self.assertIn("Lead not found", list_response.json()["detail"])

    def test_career_context_round_trip(self) -> None:
        get_response = self.client.get("/personal/career-context")
        self.assertEqual(get_response.status_code, 200)
        self.assertEqual(get_response.json()["current_headline"], "")

        put_response = self.client.put(
            "/personal/career-context",
            json={
                "current_headline": "AI Product Engineer",
                "target_roles": ["Founding Engineer", "AI Engineer"],
                "core_skills": ["python", "fastapi"],
                "technical_stack": ["next.js", "typescript"],
                "project_highlights": ["Built local-first opportunity pipeline"],
                "industries": ["SaaS"],
                "location_preferences": ["Remote"],
                "visa_notes": "No sponsorship required",
                "preferred_opportunity_types": ["job", "client"],
                "positioning_statement": "Bridge AI prototypes to reliable product workflows.",
                "proof_points": ["Reduced manual review time by 40%"],
                "raw_cv_text": "10+ years building product systems",
            },
        )
        self.assertEqual(put_response.status_code, 200)
        self.assertEqual(put_response.json()["current_headline"], "AI Product Engineer")

        fetch_response = self.client.get("/personal/career-context")
        self.assertEqual(fetch_response.status_code, 200)
        self.assertEqual(fetch_response.json()["core_skills"], ["python", "fastapi"])

    def test_career_context_extract_from_pasted_text(self) -> None:
        response = self.client.post(
            "/personal/career-context/extract",
            data={
                "raw_cv_text": "AI Product Engineer\nBuilt Python and FastAPI services for SaaS workflow automation.\nReduced processing time by 35%.",
                "extraction_mode": "local",
            },
        )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertIn("suggested_career_context", body)
        self.assertIn("extracted_full_name", body)
        self.assertIn("core_skills", body["suggested_career_context"])
        self.assertIn("reasoning_summary", body)

    def test_career_context_extract_quality_on_representative_cv(self) -> None:
        response = self.client.post(
            "/personal/career-context/extract",
            data={
                "raw_cv_text": _read_fixture("messy_cv_extracted_text.txt"),
                "extraction_mode": "local",
            },
        )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        suggested = body["suggested_career_context"]

        self.assertEqual(body["extracted_full_name"], "Danaishe Mamvura")
        self.assertIn("AI Product", suggested["current_headline"])
        self.assertIn("Systems Engineer", suggested["current_headline"])
        lowered_headline = suggested["current_headline"].lower()
        self.assertNotIn("berlin", lowered_headline)
        self.assertNotIn("germany", lowered_headline)
        self.assertNotIn("@", lowered_headline)
        self.assertNotIn("176", lowered_headline)
        self.assertGreater(len(suggested["target_roles"]), 0)
        self.assertTrue(
            any(
                role in {item.lower() for item in suggested["target_roles"]}
                for role in {"ai product engineer", "ai systems engineer"}
            )
        )

        proof_points = suggested["proof_points"]
        self.assertGreaterEqual(len(proof_points), 5)
        self.assertLessEqual(len(proof_points), 8)
        self.assertEqual(len({item.lower() for item in proof_points}), len(proof_points))

        proof_text = " ".join(proof_points).lower()
        self.assertIn("ai-powered football analytics system", proof_text)
        self.assertIn("ai-powered medical diagnosis & treatment assistant", proof_text)
        self.assertIn("ai-powered financial portfolio assistant", proof_text)
        self.assertIn("kindezi world", proof_text)
        self.assertIn("zim cyber city", proof_text)

        bad_labels = {
            "customtkinter:",
            "docker:",
            "present:",
            "2024:",
            "2020:",
        }
        for point in proof_points:
            lowered_point = point.lower()
            for bad in bad_labels:
                self.assertNotIn(bad, lowered_point)

        stack_lower = {item.lower() for item in suggested["technical_stack"]}
        expected_stack = {
            "python",
            "fastapi",
            "next.js",
            "react",
            "typescript",
            "docker",
            "tensorflow",
            "opencv",
            "scikit-learn",
            "langchain",
            "gpt-4 api",
            "tailwind css",
            "flutter",
            "firebase",
            "pandas",
            "numpy",
            "chart.js",
            "render",
            "vercel",
            "railway",
            "github",
        }
        self.assertGreaterEqual(len(stack_lower.intersection(expected_stack)), 12)
        for expected in {"fastapi", "next.js", "python", "typescript", "docker", "react"}:
            self.assertIn(expected, stack_lower)

        persisted_state = self.client.get("/personal/career-context")
        self.assertEqual(persisted_state.status_code, 200)
        self.assertEqual(persisted_state.json()["current_headline"], "")

    def test_career_context_extract_rejects_empty_text(self) -> None:
        response = self.client.post(
            "/personal/career-context/extract",
            data={
                "raw_cv_text": "",
                "extraction_mode": "local",
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("CV text is required", response.json()["detail"])

    def test_career_context_extract_is_draft_only_until_saved(self) -> None:
        before = self.client.get("/personal/career-context")
        self.assertEqual(before.status_code, 200)
        before_headline = before.json()["current_headline"]

        extract_response = self.client.post(
            "/personal/career-context/extract",
            data={
                "raw_cv_text": "Senior Platform Engineer\nBuilt FastAPI and TypeScript systems.",
                "extraction_mode": "local",
            },
        )
        self.assertEqual(extract_response.status_code, 200)

        after = self.client.get("/personal/career-context")
        self.assertEqual(after.status_code, 200)
        self.assertEqual(after.json()["current_headline"], before_headline)

    def test_career_context_extract_supports_txt_upload(self) -> None:
        response = self.client.post(
            "/personal/career-context/extract",
            data={"extraction_mode": "local"},
            files={
                "cv_file": (
                    "cv.txt",
                    "Data Engineer\nBuilt analytics pipeline in Python and SQL.",
                    "text/plain",
                )
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("suggested_career_context", response.json())

    def test_career_context_extract_rejects_unsupported_file_type(self) -> None:
        response = self.client.post(
            "/personal/career-context/extract",
            data={"extraction_mode": "local"},
            files={"cv_file": ("cv.md", "# my cv", "text/markdown")},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Unsupported file type", response.json()["detail"])

    def test_career_context_extract_rejects_empty_extracted_file_text(self) -> None:
        response = self.client.post(
            "/personal/career-context/extract",
            data={"extraction_mode": "local"},
            files={"cv_file": ("cv.txt", "   \n\n   ", "text/plain")},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Extracted CV text is empty", response.json()["detail"])

    def test_career_context_extract_handles_corrupt_docx(self) -> None:
        response = self.client.post(
            "/personal/career-context/extract",
            data={"extraction_mode": "local"},
            files={
                "cv_file": (
                    "cv.docx",
                    b"not-a-valid-docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Could not read DOCX file", response.json()["detail"])

    def test_career_context_extract_supports_docx_upload(self) -> None:
        doc = Document()
        doc.add_paragraph("Senior AI Engineer")
        doc.add_paragraph("Built FastAPI and TypeScript systems for SaaS workflows.")
        doc.add_paragraph("Reduced manual process time by 30%.")
        buf = BytesIO()
        doc.save(buf)
        response = self.client.post(
            "/personal/career-context/extract",
            data={"extraction_mode": "local"},
            files={
                "cv_file": (
                    "cv.docx",
                    buf.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("suggested_career_context", response.json())

    @patch("app.services.career_context_extractor._extract_pdf_text", return_value="AI Engineer\nBuilt Python workflows")
    def test_career_context_extract_supports_pdf_upload_path(self, _mock_pdf_extract) -> None:
        response = self.client.post(
            "/personal/career-context/extract",
            data={"extraction_mode": "local"},
            files={"cv_file": ("cv.pdf", b"%PDF-1.4 mock", "application/pdf")},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("suggested_career_context", response.json())

    def test_extract_from_text_requires_raw_text(self) -> None:
        response = self.client.post(
            "/personal/leads/extract-from-text",
            json={
                "raw_text": "",
                "source_type": "job_listing",
                "user_goal": "job",
                "use_career_context": True,
            },
        )
        self.assertEqual(response.status_code, 422)

    def test_extract_from_text_rejects_unsupported_source_type(self) -> None:
        response = self.client.post(
            "/personal/leads/extract-from-text",
            json={
                "raw_text": "Hiring for a backend engineer role.",
                "source_type": "sms_message",
                "user_goal": "job",
                "use_career_context": True,
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Unsupported sourceType", response.json()["detail"])

    def test_extract_from_text_returns_warning_with_empty_context(self) -> None:
        response = self.client.post(
            "/personal/leads/extract-from-text",
            json={
                "raw_text": "We are hiring a python and fastapi engineer to build automation workflows.",
                "source_type": "job_listing",
                "user_goal": "job",
                "use_career_context": True,
            },
        )
        self.assertEqual(response.status_code, 200)
        warnings = response.json()["review_warnings"]
        self.assertTrue(any("Career context is empty" in item for item in warnings))

    def test_extract_from_text_is_draft_only_and_can_be_saved_via_normal_flow(self) -> None:
        extract_response = self.client.post(
            "/personal/leads/extract-from-text",
            json={
                "raw_text": "Recruiter message: looking for a senior fastapi engineer. Contact: Alex Rivera. Company: Orbital Labs.",
                "source_type": "recruiter_message",
                "user_goal": "job",
                "use_career_context": False,
            },
        )
        self.assertEqual(extract_response.status_code, 200)
        payload = extract_response.json()
        self.assertIn("career_fit_score", payload)
        self.assertIn("reasoning_summary", payload)

        leads_response = self.client.get("/personal/leads")
        self.assertEqual(leads_response.status_code, 200)
        self.assertEqual(len(leads_response.json()), 0)

        rule_actions_before = self.client.get("/rule-of-100/actions")
        self.assertEqual(rule_actions_before.status_code, 200)
        before_count = len(rule_actions_before.json())

        suggested = payload["suggested_lead"]
        create_response = self.client.post(
            "/personal/leads",
            json={
                "name": suggested["name"],
                "role": suggested["role"],
                "organisation": suggested["organisation"],
                "organisation_website": suggested["organisation_website"],
                "linkedin_url": suggested["linkedin_url"],
                "email": suggested["email"],
                "location": suggested["location"],
                "source": suggested["source"],
                "opportunity_type": suggested["opportunity_type"],
                "relationship_strength": suggested["relationship_strength"],
                "status": "new",
                "fit_score": suggested["fit_score"],
                "priority": suggested["priority"],
                "problem_observed": suggested["problem_observed"],
                "why_relevant": suggested["why_relevant"],
                "suggested_angle": suggested["suggested_angle"],
                "notes": suggested["notes"],
                "tags": suggested["tags"],
                "next_action": suggested["next_action"],
                "follow_up_date": suggested["follow_up_date"],
            },
        )
        self.assertEqual(create_response.status_code, 200)

        leads_after = self.client.get("/personal/leads")
        self.assertEqual(leads_after.status_code, 200)
        self.assertEqual(len(leads_after.json()), 1)

        rule_actions_after = self.client.get("/rule-of-100/actions")
        self.assertEqual(rule_actions_after.status_code, 200)
        self.assertEqual(len(rule_actions_after.json()), before_count)

    def test_extract_from_linkedin_job_listing_quality(self) -> None:
        self.client.put(
            "/personal/career-context",
            json={
                "current_headline": "AI Product & Systems Engineer",
                "target_roles": ["Software Engineer", "AI Engineer"],
                "core_skills": ["python", "fastapi", "react", "typescript", "llm api"],
                "technical_stack": ["rag", "langchain", "docker", "firebase", "flutter", "chart.js"],
                "project_highlights": ["Built AI feature delivery systems"],
                "industries": ["GovTech"],
                "location_preferences": ["Berlin"],
                "visa_notes": "",
                "preferred_opportunity_types": ["job"],
                "positioning_statement": "",
                "proof_points": ["Shipped production AI-assisted workflows"],
                "raw_cv_text": "Built Python/FastAPI and React/TypeScript systems with LLM integrations.",
            },
        )

        response = self.client.post(
            "/personal/leads/extract-from-text",
            json={
                "raw_text": _read_fixture("linkedin_job_listing_admi_kommunal.txt"),
                "source_type": "job_listing",
                "optional_source_url": "https://www.linkedin.com/jobs/view/123456",
                "user_goal": "job",
                "use_career_context": True,
            },
        )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        lead = body["suggested_lead"]

        self.assertIn(lead["organisation"], {"admi Kommunal", "admi Kommunal GmbH"})
        self.assertNotEqual(lead["organisation"].lower(), "about the job")
        self.assertEqual(lead["role"], "Software Engineer - Fullstack & AI (m/w/d)")
        self.assertIn("Berlin", lead["location"])
        self.assertIn("Germany", lead["location"])
        self.assertNotEqual(lead["location"].strip().lower(), "remote")
        self.assertEqual(lead["opportunity_type"], "job")
        self.assertEqual(lead["source"], "capture:job_listing")
        self.assertEqual(lead["name"], "")

        combined = " ".join([lead["name"], lead["role"], lead["organisation"]]).lower()
        self.assertNotIn("premium", combined)
        self.assertNotIn("try premium", combined)

        missing_fields = set(body["missing_fields"])
        self.assertNotIn("name", missing_fields)

        matched = {item.lower() for item in body["matched_skills"]}
        self.assertIn("python", matched)
        self.assertIn("fastapi", matched)
        self.assertTrue("react" in matched or "typescript" in matched)
        self.assertIn("llm apis", matched)

        gaps = {item.lower() for item in body["missing_skills_or_gaps"]}
        self.assertTrue("kotlin/jvm" in gaps or "spring boot" in gaps)
        self.assertTrue("postgresql" in gaps or "gcp" in gaps)
        self.assertIn("langgraph", gaps)

    def test_create_rule_action_for_job_lead_uses_application_tailoring(self) -> None:
        create_response = self.client.post(
            "/personal/leads",
            json={
                "name": "",
                "role": "Software Engineer - Fullstack & AI (m/w/d)",
                "organisation": "admi Kommunal",
                "source": "capture:job_listing",
                "opportunity_type": "job",
                "relationship_strength": "cold",
                "status": "new",
                "fit_score": 7,
                "priority": "medium",
                "problem_observed": "Needs Kotlin/Spring ramp-up framing",
                "why_relevant": "Strong Python/FastAPI and React/TypeScript overlap with role requirements",
                "suggested_angle": "Emphasize AI/full-stack product implementation",
                "notes": "Job metadata: work mode: on-site, employment: full-time",
                "tags": ["python", "fastapi", "react", "typescript", "llm apis"],
                "next_action": "review and tailor application",
            },
        )
        self.assertEqual(create_response.status_code, 200)
        lead_id = create_response.json()["id"]

        action_response = self.client.post(f"/personal/leads/{lead_id}/create-rule-action")
        self.assertEqual(action_response.status_code, 200)
        action = action_response.json()

        self.assertEqual(action["channel"], "job_application")
        self.assertIn("application", action["action_type"].lower())
        self.assertIn("tailor application", action["suggested_action"].lower())
        self.assertIn("software engineer - fullstack & ai", action["suggested_action"].lower())
        self.assertIn("admi kommunal", action["suggested_action"].lower())

        message = action["suggested_message"]
        self.assertFalse(message.lower().startswith("hi,"))
        self.assertIn("review the requirements", message.lower())
        self.assertIn("python/fastapi", message.lower())
        self.assertIn("react/typescript", message.lower())
        self.assertIn("llm/rag", message.lower())
        self.assertIn("kotlin/spring", message.lower())

    def test_duplicate_lead_creation_is_blocked_for_same_source_org_role(self) -> None:
        payload = {
            "name": "",
            "role": "Software Engineer - Fullstack & AI (m/w/d)",
            "organisation": "admi Kommunal",
            "organisation_website": "",
            "linkedin_url": "https://www.linkedin.com/jobs/view/123456",
            "email": "",
            "location": "Berlin, Germany",
            "source": "capture:job_listing",
            "opportunity_type": "job",
            "relationship_strength": "cold",
            "status": "new",
            "fit_score": 7,
            "priority": "medium",
            "problem_observed": "",
            "why_relevant": "",
            "suggested_angle": "",
            "notes": "",
            "tags": ["python"],
            "next_action": "review and tailor application",
            "follow_up_date": None,
        }

        first = self.client.post("/personal/leads", json=payload)
        self.assertEqual(first.status_code, 200)

        second = self.client.post("/personal/leads", json=payload)
        self.assertEqual(second.status_code, 400)
        self.assertIn("Existing lead found", second.json()["detail"])

    def test_deleting_one_lead_only_removes_that_lead(self) -> None:
        first = self.client.post(
            "/personal/leads",
            json={
                "name": "",
                "role": "Role A",
                "organisation": "Org A",
                "source": "manual",
                "opportunity_type": "job",
                "relationship_strength": "cold",
                "status": "new",
                "fit_score": 5,
                "priority": "medium",
                "next_action": "review",
            },
        )
        second = self.client.post(
            "/personal/leads",
            json={
                "name": "",
                "role": "Role B",
                "organisation": "Org B",
                "source": "manual",
                "opportunity_type": "job",
                "relationship_strength": "cold",
                "status": "new",
                "fit_score": 5,
                "priority": "medium",
                "next_action": "review",
            },
        )
        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)

        first_id = first.json()["id"]
        second_id = second.json()["id"]

        delete_response = self.client.delete(f"/personal/leads/{first_id}")
        self.assertEqual(delete_response.status_code, 200)
        self.assertTrue(delete_response.json()["deleted"])

        list_response = self.client.get("/personal/leads")
        self.assertEqual(list_response.status_code, 200)
        ids = {item["id"] for item in list_response.json()}
        self.assertNotIn(first_id, ids)
        self.assertIn(second_id, ids)


if __name__ == "__main__":
    unittest.main()
